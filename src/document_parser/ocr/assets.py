from __future__ import annotations

import base64
from dataclasses import dataclass
from pathlib import Path
import shutil
from typing import Any
from urllib.parse import unquote_to_bytes, urlparse
from zipfile import ZipFile

from ..models import Block, ParsedDocument


@dataclass
class MaterializedAsset:
    status: str
    path: str = ""
    reason_code: str = ""
    reason: str = ""
    metadata: dict[str, Any] | None = None


class ImageAssetMaterializer:
    """Extract an image Block into a local, auditable file for OCR."""

    def materialize(self, document: ParsedDocument, block: Block, output_dir: Path) -> MaterializedAsset:
        output_dir.mkdir(parents=True, exist_ok=True)
        handler = getattr(self, f"_from_{document.file_type}", None)
        if handler is None:
            return MaterializedAsset(
                status="unavailable",
                reason_code="UNSUPPORTED_IMAGE_SOURCE",
                reason=f"尚未实现 {document.file_type} 图片资源提取。",
            )
        try:
            return handler(document, block, output_dir)
        except Exception as error:
            return MaterializedAsset(
                status="failed",
                reason_code="IMAGE_MATERIALIZATION_FAILED",
                reason=f"图片资源提取失败：{type(error).__name__}。",
            )

    def _from_docx(
        self, document: ParsedDocument, block: Block, output_dir: Path
    ) -> MaterializedAsset:
        member = str(block.image.image_path or "").lstrip("/") if block.image else ""
        if not member:
            return _missing_asset("DOCX_IMAGE_PART_MISSING", "DOCX 图片没有 relationship part 路径。")
        with ZipFile(document.source_path) as archive:
            data = archive.read(member)
        path = output_dir / f"{block.block_id}{Path(member).suffix or '.bin'}"
        path.write_bytes(data)
        return MaterializedAsset("ready", str(path), metadata={"archive_member": member})

    def _from_pptx(
        self, document: ParsedDocument, block: Block, output_dir: Path
    ) -> MaterializedAsset:
        from pptx import Presentation

        trace = block.source_trace
        if trace is None or trace.raw_slide_index is None or trace.raw_object_index is None:
            return _missing_asset("PPTX_PICTURE_LOCATOR_MISSING", "PPTX 图片缺少 slide/shape 定位信息。")
        presentation = Presentation(document.source_path)
        slide = presentation.slides[trace.raw_slide_index]
        shapes = sorted(
            slide.shapes,
            key=lambda shape: (getattr(shape, "top", 0), getattr(shape, "left", 0)),
        )
        shape = shapes[trace.raw_object_index]
        image = shape.image
        path = output_dir / f"{block.block_id}.{image.ext or 'bin'}"
        path.write_bytes(image.blob)
        return MaterializedAsset("ready", str(path))

    def _from_xlsx(
        self, document: ParsedDocument, block: Block, output_dir: Path
    ) -> MaterializedAsset:
        from openpyxl import load_workbook

        trace = block.source_trace
        image_index = int(block.metadata.get("image_index", -1))
        if trace is None or not trace.raw_sheet_name or image_index < 0:
            return _missing_asset("XLSX_IMAGE_LOCATOR_MISSING", "XLSX 图片缺少 Sheet 或索引。")
        workbook = load_workbook(document.source_path, data_only=False)
        try:
            image = list(getattr(workbook[trace.raw_sheet_name], "_images", []) or [])[image_index]
            data = image._data()
            extension = str(getattr(image, "format", "") or "png").lower()
        finally:
            workbook.close()
        path = output_dir / f"{block.block_id}.{extension}"
        path.write_bytes(data)
        return MaterializedAsset("ready", str(path))

    def _from_html(
        self, document: ParsedDocument, block: Block, output_dir: Path
    ) -> MaterializedAsset:
        return self._from_linked_asset(document, block, output_dir)

    def _from_markdown(
        self, document: ParsedDocument, block: Block, output_dir: Path
    ) -> MaterializedAsset:
        return self._from_linked_asset(document, block, output_dir)

    def _from_linked_asset(
        self, document: ParsedDocument, block: Block, output_dir: Path
    ) -> MaterializedAsset:
        value = str(block.image.image_path or "").strip() if block.image else ""
        if not value:
            return _missing_asset("IMAGE_REFERENCE_MISSING", "图片引用为空。")
        if value.startswith("data:"):
            return self._from_data_uri(value, block, output_dir)
        parsed = urlparse(value)
        if parsed.scheme in {"http", "https"}:
            return _missing_asset(
                "REMOTE_IMAGE_REQUIRES_FETCH_POLICY",
                "远程图片未自动下载，需要先经过域名白名单和下载策略。",
            )
        if parsed.scheme and parsed.scheme != "file":
            return _missing_asset("UNSUPPORTED_IMAGE_URI", f"不支持的图片 URI：{parsed.scheme}。")
        source_dir = Path(document.source_path).resolve().parent
        candidate = Path(parsed.path)
        candidate = candidate.resolve() if candidate.is_absolute() else (source_dir / candidate).resolve()
        if not candidate.is_relative_to(source_dir):
            return _missing_asset(
                "IMAGE_PATH_OUTSIDE_SOURCE_ROOT",
                "图片路径越过源文档目录，安全策略拒绝自动读取。",
            )
        if not candidate.is_file():
            return _missing_asset("LOCAL_IMAGE_NOT_FOUND", "本地图片文件不存在。")
        path = output_dir / f"{block.block_id}{candidate.suffix or '.bin'}"
        shutil.copyfile(candidate, path)
        return MaterializedAsset("ready", str(path), metadata={"source_asset_path": str(candidate)})

    def _from_data_uri(self, value: str, block: Block, output_dir: Path) -> MaterializedAsset:
        header, payload = value.split(",", 1)
        media_type = header[5:].split(";", 1)[0]
        extension = {
            "image/png": ".png",
            "image/jpeg": ".jpg",
            "image/webp": ".webp",
            "image/gif": ".gif",
        }.get(media_type, ".bin")
        data = base64.b64decode(payload) if ";base64" in header else unquote_to_bytes(payload)
        path = output_dir / f"{block.block_id}{extension}"
        path.write_bytes(data)
        return MaterializedAsset("ready", str(path), metadata={"media_type": media_type})


def _missing_asset(code: str, reason: str) -> MaterializedAsset:
    return MaterializedAsset(status="unavailable", reason_code=code, reason=reason)
