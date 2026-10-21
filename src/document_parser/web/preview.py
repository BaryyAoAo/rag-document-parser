"""Small, read-only Office previews without an external document server."""

from __future__ import annotations

from html import escape
from pathlib import Path


def _table(rows: list[list[str]]) -> str:
    body = []
    for row in rows[:100]:
        body.append("<tr>" + "".join(f"<td>{escape(str(cell))}</td>" for cell in row[:30]) + "</tr>")
    return "<div class='table-wrap'><table>" + "".join(body) + "</table></div>"


def office_preview_html(path: Path) -> str:
    suffix = path.suffix.lower()
    sections: list[str] = []
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets[:20]:
                rows = [["" if cell is None else str(cell) for cell in row[:30]]
                        for row in sheet.iter_rows(max_row=100, max_col=30, values_only=True)]
                sections.append(f"<section><h2>{escape(sheet.title)}</h2>{_table(rows)}</section>")
        finally:
            book.close()
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(path)
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph) and item.text.strip():
                style = item.style.name.lower() if item.style else ""
                tag = "h2" if "heading" in style or "标题" in style else "p"
                sections.append(f"<{tag}>{escape(item.text)}</{tag}>")
            elif isinstance(item, Table):
                sections.append(_table([[cell.text for cell in row.cells] for row in item.rows]))
    elif suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        for index, slide in enumerate(presentation.slides, 1):
            content = [f"<h2>第 {index} 页</h2>"]
            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    content.append(_table([[cell.text for cell in row.cells] for row in shape.table.rows]))
                elif getattr(shape, "has_text_frame", False) and shape.text.strip():
                    content.append(f"<p>{escape(shape.text)}</p>")
                elif getattr(shape, "shape_type", None) == 13:
                    content.append("<p class='muted'>[图片]</p>")
            sections.append("<section class='slide'>" + "".join(content) + "</section>")
    else:
        raise ValueError("Office preview supports DOCX, XLSX and PPTX")
    body = "".join(sections) or "<p>文件中没有可显示的文字或表格。</p>"
    return (
        "<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'>"
        "<meta name='viewport' content='width=device-width,initial-scale=1'>"
        "<style>body{font:14px/1.65 system-ui,sans-serif;color:#26382c;background:#fff;margin:0;padding:18px}"
        "h1{font-size:18px}h2{font-size:15px;margin:0 0 12px}section{border:1px solid #e4eae2;"
        "border-radius:8px;padding:16px;margin:0 0 16px}p{white-space:pre-wrap;margin:8px 0}"
        "table{border-collapse:collapse;min-width:100%;font-size:12px}td{border:1px solid #d5ddd6;"
        "padding:6px;white-space:pre-wrap;vertical-align:top}.table-wrap{overflow:auto;max-height:620px}"
        ".slide{min-height:210px}.muted{color:#93a296}</style></head><body>"
        f"<h1>{escape(path.name)}</h1>{body}</body></html>"
    )
