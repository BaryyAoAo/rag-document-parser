from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ..models import Block, ImageData, ParserResult, SourceFile, SourceTrace
from ..quality import build_quality_report
from ..text_builders import normalize_text, table_from_rows
from .base import BaseParser


class XlsxParser(BaseParser):
    parser_name = "openpyxl"
    file_type = "xlsx"

    def parse(self, source: SourceFile) -> ParserResult:
        path = Path(source.file_path)
        document_id = self.build_document_id(path)
        metadata = self.base_metadata(source, path)
        workbook = load_workbook(path, data_only=False)
        metadata["sheet_names"] = workbook.sheetnames
        title = path.stem
        raw_outputs = []
        blocks: list[Block] = []
        order = 0

        for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cleaned_row = [normalize_text(cell) for cell in row]
                if any(cleaned_row):
                    rows.append(cleaned_row)

            raw_outputs.append(
                {
                    "sheet_index": sheet_index,
                    "sheet_name": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                    "image_count": len(getattr(sheet, "_images", []) or []),
                    "rows_sample": rows[:10],
                }
            )

            if rows:
                table_data = table_from_rows(rows, caption=f"工作表：{sheet.title}")
                order += 1
                blocks.append(
                    Block(
                        block_id=f"{document_id}_sheet_{sheet_index:04d}_tbl_0001",
                        document_id=document_id,
                        block_type="table",
                        table_data=table_data,
                        metadata={
                            "source_type": "xlsx",
                            "sheet_name": sheet.title,
                            "sheet_index": sheet_index,
                            "cell_range": f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}",
                            "table_id": f"sheet_{sheet_index:04d}_table_0001",
                            "row_count": len(table_data.rows),
                            "column_count": len(table_data.headers),
                        },
                        source_trace=SourceTrace(
                            parser=self.parser_name,
                            raw_object_type="worksheet_table",
                            raw_sheet_name=sheet.title,
                            raw_object_index=0,
                        ),
                        order=order,
                    )
                )

            for image_index, image in enumerate(getattr(sheet, "_images", []) or []):
                order += 1
                anchor = getattr(image, "anchor", None)
                anchor_from = getattr(anchor, "_from", None)
                row_index = int(getattr(anchor_from, "row", 0)) + 1 if anchor_from else None
                column_index = int(getattr(anchor_from, "col", 0)) + 1 if anchor_from else None
                blocks.append(
                    Block(
                        block_id=f"{document_id}_sheet_{sheet_index:04d}_img_{image_index + 1:04d}",
                        document_id=document_id,
                        block_type="image",
                        image=ImageData(
                            image_id=f"sheet_{sheet_index:04d}_img_{image_index + 1:04d}",
                            image_path=f"xlsx://{sheet.title}/{image_index}",
                            width=float(getattr(image, "width", 0) or 0) or None,
                            height=float(getattr(image, "height", 0) or 0) or None,
                        ),
                        metadata={
                            "source_type": "xlsx",
                            "sheet_name": sheet.title,
                            "sheet_index": sheet_index,
                            "image_index": image_index,
                            "anchor_row": row_index,
                            "anchor_column": column_index,
                            "ocr_decision": "review_before_ocr",
                            "ocr_status": "pending_review",
                            "ocr_reason": "XLSX 工作表图片可能包含截图或扫描表格",
                            "ocr_reason_codes": ["XLSX_EMBEDDED_IMAGE"],
                        },
                        source_trace=SourceTrace(
                            parser=self.parser_name,
                            raw_object_type="worksheet_image",
                            raw_sheet_name=sheet.title,
                            raw_object_index=image_index,
                        ),
                        quality_flags=["review_before_ocr"],
                        order=order,
                    )
                )

        workbook.close()
        quality_report = build_quality_report(document_id, raw_outputs, blocks)
        return ParserResult(
            document_id=document_id,
            title=title,
            file_type=self.file_type,
            source_path=str(path),
            source_url=source.source_url,
            metadata=metadata,
            raw_outputs=raw_outputs,
            blocks=blocks,
            quality_report=quality_report,
        )
