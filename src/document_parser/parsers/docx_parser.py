from __future__ import annotations

from pathlib import Path
from typing import Iterator
from zipfile import ZipFile

from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..models import Block, ImageData, ParserResult, SourceFile, SourceTrace
from ..quality import build_quality_report
from ..text_builders import normalize_text, table_from_rows
from .base import BaseParser


class DocxParser(BaseParser):
    parser_name = "python-docx"
    file_type = "docx"

    def parse(self, source: SourceFile) -> ParserResult:
        path = Path(source.file_path)
        document_id = self.build_document_id(path)
        metadata = self.base_metadata(source, path)
        doc = DocxDocument(path)
        title = self._guess_title(doc, path)
        image_names = self._list_media_names(path)
        metadata["image_count"] = len(image_names)

        raw_outputs = [
            {
                "document_body_order": [],
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
                "media_files": image_names,
            }
        ]
        blocks: list[Block] = []
        order = 0

        for item_index, item in enumerate(self._iter_body_items(doc), start=1):
            if isinstance(item, Paragraph):
                text = normalize_text(item.text)
                style_name = item.style.name if item.style else ""
                image_rids = self._paragraph_image_rids(item)
                raw_outputs[0]["document_body_order"].append(
                    {
                        "index": item_index,
                        "type": "paragraph",
                        "style": style_name,
                        "text": text,
                        "image_rids": image_rids,
                    }
                )
                if text:
                    order += 1
                    block_type, heading_level = self._paragraph_block_type(style_name, text)
                    blocks.append(
                        Block(
                            block_id=f"{document_id}_docx_txt_{order:04d}",
                            document_id=document_id,
                            block_type=block_type,
                            text=text,
                            metadata={
                                "source_type": "docx",
                                "style_name": style_name,
                                "heading_level": heading_level,
                            },
                            source_trace=SourceTrace(
                                parser=self.parser_name,
                                raw_object_type="paragraph",
                                raw_object_index=item_index - 1,
                            ),
                            order=order,
                        )
                    )
                for rid_index, rid in enumerate(image_rids, start=1):
                    order += 1
                    blocks.append(
                        Block(
                            block_id=f"{document_id}_docx_img_{order:04d}",
                            document_id=document_id,
                            block_type="image",
                            image=ImageData(image_id=f"img_{order:04d}", image_path=self._image_part_name(doc, rid)),
                            metadata={
                                "source_type": "docx",
                                "relationship_id": rid,
                                "ocr_decision": "review_before_ocr",
                                "ocr_status": "pending_review",
                                "ocr_reason": "docx paragraph contains embedded image",
                                "ocr_reason_codes": ["DOCX_EMBEDDED_IMAGE"],
                            },
                            source_trace=SourceTrace(
                                parser=self.parser_name,
                                raw_object_type="paragraph_image",
                                raw_object_index=rid_index - 1,
                                raw_locator=rid,
                            ),
                            quality_flags=["review_before_ocr"],
                            order=order,
                        )
                    )
            elif isinstance(item, Table):
                rows = [[cell.text for cell in row.cells] for row in item.rows]
                raw_outputs[0]["document_body_order"].append(
                    {
                        "index": item_index,
                        "type": "table",
                        "row_count": len(rows),
                        "rows_sample": rows[:5],
                    }
                )
                table_data = table_from_rows(rows)
                if table_data.headers or table_data.rows:
                    order += 1
                    blocks.append(
                        Block(
                            block_id=f"{document_id}_docx_tbl_{order:04d}",
                            document_id=document_id,
                            block_type="table",
                            table_data=table_data,
                            metadata={
                                "source_type": "docx",
                                "table_id": f"tbl_{order:04d}",
                                "row_count": len(table_data.rows),
                                "column_count": len(table_data.headers),
                            },
                            source_trace=SourceTrace(
                                parser=self.parser_name,
                                raw_object_type="table",
                                raw_object_index=item_index - 1,
                            ),
                            order=order,
                        )
                    )

        quality_report = build_quality_report(document_id, raw_outputs, blocks)
        return ParserResult(
            document_id=document_id,
            title=title,
            file_type=self.file_type,
            source_path=str(path),
            source_url=source.source_url,
            metadata=metadata,
            raw_outputs=raw_outputs,
            blocks=blocks,
            quality_report=quality_report,
        )

    def _iter_body_items(self, doc: DocxDocumentType) -> Iterator[Paragraph | Table]:
        for child in doc.element.body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, doc)
            elif isinstance(child, CT_Tbl):
                yield Table(child, doc)

    def _guess_title(self, doc: DocxDocumentType, path: Path) -> str:
        for paragraph in doc.paragraphs:
            text = normalize_text(paragraph.text)
            if text:
                return text[:100]
        return path.stem

    def _list_media_names(self, path: Path) -> list[str]:
        with ZipFile(path) as archive:
            return [name for name in archive.namelist() if name.startswith("word/media/")]

    def _paragraph_block_type(self, style_name: str, text: str) -> tuple[str, int | None]:
        if style_name.startswith("Heading") or style_name.startswith("标题"):
            digits = "".join(ch for ch in style_name if ch.isdigit())
            return "heading", int(digits) if digits else 1
        if style_name == "Title" or len(text) <= 40 and text[:2] in {"一、", "二、", "三、", "四、"}:
            return "heading", 1
        return "paragraph", None

    def _paragraph_image_rids(self, paragraph: Paragraph) -> list[str]:
        rids: list[str] = []
        rel_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        for element in paragraph._element.iter():
            if element.tag.endswith("}blip"):
                rid = element.attrib.get(rel_attr)
                if rid:
                    rids.append(rid)
        return rids

    def _image_part_name(self, doc: DocxDocumentType, rid: str) -> str | None:
        rel = doc.part.rels.get(rid)
        if not rel:
            return None
        target_part = getattr(rel, "target_part", None)
        partname = getattr(target_part, "partname", None)
        return str(partname) if partname else None
