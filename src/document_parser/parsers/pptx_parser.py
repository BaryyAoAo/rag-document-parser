from __future__ import annotations

from pathlib import Path

from ..models import Block, ImageData, ParserResult, QualityWarning, SourceFile, SourceTrace
from ..quality import build_quality_report
from ..text_builders import normalize_text, table_from_rows
from .base import BaseParser


class PptxParser(BaseParser):
    parser_name = "python-pptx"
    file_type = "pptx"

    def parse(self, source: SourceFile) -> ParserResult:
        path = Path(source.file_path)
        document_id = self.build_document_id(path)
        metadata = self.base_metadata(source, path)
        try:
            from pptx import Presentation
        except Exception as exc:  # pragma: no cover - depends on optional dependency
            warning = QualityWarning(
                level="error",
                warning_type="missing_dependency",
                message=f"python-pptx is not available: {exc}",
            )
            quality_report = build_quality_report(document_id, [], [], [warning])
            return ParserResult(
                document_id=document_id,
                title=path.stem,
                file_type=self.file_type,
                source_path=str(path),
                source_url=source.source_url,
                metadata=metadata,
                raw_outputs=[],
                blocks=[],
                quality_report=quality_report,
            )

        presentation = Presentation(path)
        metadata["slide_count"] = len(presentation.slides)
        raw_outputs = []
        blocks: list[Block] = []
        order = 0

        for slide_index, slide in enumerate(presentation.slides, start=1):
            raw_slide = {"slide_index": slide_index, "shape_count": len(slide.shapes), "shapes": []}
            shapes = sorted(slide.shapes, key=lambda shape: (getattr(shape, "top", 0), getattr(shape, "left", 0)))
            for shape_index, shape in enumerate(shapes, start=1):
                raw_slide["shapes"].append(
                    {
                        "shape_index": shape_index,
                        "shape_type": str(getattr(shape, "shape_type", "")),
                        "has_text_frame": getattr(shape, "has_text_frame", False),
                        "has_table": getattr(shape, "has_table", False),
                    }
                )
                if getattr(shape, "has_table", False):
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    table_data = table_from_rows(rows)
                    order += 1
                    blocks.append(
                        Block(
                            block_id=f"{document_id}_slide_{slide_index:04d}_tbl_{order:04d}",
                            document_id=document_id,
                            block_type="table",
                            table_data=table_data,
                            page_start=slide_index,
                            page_end=slide_index,
                            metadata={"source_type": "pptx", "slide_index": slide_index, "table_id": f"tbl_{order:04d}"},
                            source_trace=SourceTrace(
                                parser=self.parser_name,
                                raw_slide_index=slide_index - 1,
                                raw_object_type="shape_table",
                                raw_object_index=shape_index - 1,
                            ),
                            order=order,
                        )
                    )
                elif getattr(shape, "has_text_frame", False):
                    text = normalize_text(shape.text)
                    if not text:
                        continue
                    order += 1
                    blocks.append(
                        Block(
                            block_id=f"{document_id}_slide_{slide_index:04d}_txt_{order:04d}",
                            document_id=document_id,
                            block_type="heading" if shape_index == 1 and len(text) <= 60 else "paragraph",
                            text=text,
                            page_start=slide_index,
                            page_end=slide_index,
                            metadata={"source_type": "pptx", "slide_index": slide_index},
                            source_trace=SourceTrace(
                                parser=self.parser_name,
                                raw_slide_index=slide_index - 1,
                                raw_object_type="shape_text",
                                raw_object_index=shape_index - 1,
                            ),
                            order=order,
                        )
                    )
                elif "PICTURE" in str(getattr(shape, "shape_type", "")):
                    order += 1
                    blocks.append(
                        Block(
                            block_id=f"{document_id}_slide_{slide_index:04d}_img_{order:04d}",
                            document_id=document_id,
                            block_type="image",
                            image=ImageData(
                                image_id=f"img_{order:04d}",
                                width=float(getattr(shape.image, "size", (0, 0))[0] or 0) or None,
                                height=float(getattr(shape.image, "size", (0, 0))[1] or 0) or None,
                            ),
                            page_start=slide_index,
                            page_end=slide_index,
                            metadata={
                                "source_type": "pptx",
                                "slide_index": slide_index,
                                "ocr_decision": "review_before_ocr",
                                "ocr_status": "pending_review",
                                "ocr_reason": "PPTX 图片可能包含流程、截图或业务文字",
                                "ocr_reason_codes": ["PPTX_PICTURE_SHAPE"],
                            },
                            source_trace=SourceTrace(
                                parser=self.parser_name,
                                raw_slide_index=slide_index - 1,
                                raw_object_type="shape_picture",
                                raw_object_index=shape_index - 1,
                            ),
                            quality_flags=["review_before_ocr"],
                            order=order,
                        )
                    )
            raw_outputs.append(raw_slide)

        quality_report = build_quality_report(document_id, raw_outputs, blocks)
        return ParserResult(
            document_id=document_id,
            title=path.stem,
            file_type=self.file_type,
            source_path=str(path),
            source_url=source.source_url,
            metadata=metadata,
            raw_outputs=raw_outputs,
            blocks=blocks,
            quality_report=quality_report,
        )
